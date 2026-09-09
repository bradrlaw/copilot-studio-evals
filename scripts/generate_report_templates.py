#!/usr/bin/env python3
"""Generate generic, editable PowerPoint templates for agent evaluation reporting."""

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


NAVY = RGBColor(0x17, 0x2B, 0x4D)
BLUE = RGBColor(0x00, 0x66, 0xCC)
TEAL = RGBColor(0x00, 0x7A, 0x78)
GREEN = RGBColor(0x1B, 0x7F, 0x4B)
AMBER = RGBColor(0xB5, 0x68, 0x00)
RED = RGBColor(0xB4, 0x23, 0x18)
LIGHT_BLUE = RGBColor(0xE8, 0xF1, 0xFB)
LIGHT_GREY = RGBColor(0xF3, 0xF5, 0xF7)
MID_GREY = RGBColor(0x5E, 0x6C, 0x84)
DARK = RGBColor(0x17, 0x24, 0x2B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def new_deck(title: str, subject: str) -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = title
    prs.core_properties.subject = subject
    prs.core_properties.author = "Copilot Studio Evaluation Framework"
    prs.core_properties.keywords = "Copilot Studio, evaluation, template"
    prs.core_properties.comments = "Generic template generated from synthetic placeholder content."
    return prs


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 20,
    color: RGBColor = DARK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = 5 if radius else 1
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def add_header(slide, title: str, section: str):
    add_rect(slide, 0, 0, 13.333, 0.14, BLUE)
    add_text(slide, section.upper(), 0.65, 0.35, 4.0, 0.3, size=10, color=BLUE, bold=True)
    add_text(slide, title, 0.65, 0.72, 12.0, 0.65, size=28, color=NAVY, bold=True)


def add_footer(slide, number: int):
    add_text(
        slide,
        "Copilot Studio Evaluation Framework | Replace all <PLACEHOLDERS>",
        0.65,
        7.12,
        10.5,
        0.2,
        size=9,
        color=MID_GREY,
    )
    add_text(slide, str(number), 12.1, 7.08, 0.55, 0.25, size=9, color=MID_GREY, align=PP_ALIGN.RIGHT)


def add_bullets(slide, items, x, y, w, h, size=18, color=DARK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(9)
        paragraph.text = f"• {item}"
    return box


def add_title_slide(prs, title: str, subtitle: str, fields):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, 13.333, 7.5, NAVY)
    add_rect(slide, 0, 0, 0.18, 7.5, BLUE)
    add_text(slide, title, 0.85, 1.25, 11.4, 1.45, size=38, color=WHITE, bold=True)
    add_text(slide, subtitle, 0.88, 2.85, 10.8, 0.7, size=22, color=LIGHT_BLUE)
    add_text(slide, "\n".join(fields), 0.88, 4.25, 8.2, 1.3, size=16, color=WHITE)
    add_text(
        slide,
        "TEMPLATE",
        10.55,
        5.7,
        1.75,
        0.48,
        size=14,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    ).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = LIGHT_BLUE


def add_metric(slide, x, label, value, color):
    add_rect(slide, x, 1.72, 2.85, 1.35, LIGHT_GREY, color, radius=True)
    add_text(slide, value, x + 0.15, 1.93, 2.55, 0.48, size=27, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.15, 2.48, 2.55, 0.3, size=11, color=MID_GREY, bold=True, align=PP_ALIGN.CENTER)


def add_two_column_slide(prs, title, section, left_title, left_items, right_title, right_items):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, title, section)
    add_rect(slide, 0.65, 1.58, 5.85, 4.95, LIGHT_GREY, LIGHT_GREY, radius=True)
    add_rect(slide, 6.82, 1.58, 5.85, 4.95, LIGHT_BLUE, LIGHT_BLUE, radius=True)
    add_text(slide, left_title, 0.95, 1.88, 5.2, 0.45, size=18, color=NAVY, bold=True)
    add_bullets(slide, left_items, 0.95, 2.5, 5.15, 3.7, size=16)
    add_text(slide, right_title, 7.12, 1.88, 5.2, 0.45, size=18, color=NAVY, bold=True)
    add_bullets(slide, right_items, 7.12, 2.5, 5.15, 3.7, size=16)
    add_footer(slide, len(prs.slides))


def build_baseline_template(path: Path):
    prs = new_deck(
        "Agent Evaluation Baseline Report Template",
        "Generic Copilot Studio agent evaluation baseline template",
    )
    add_title_slide(
        prs,
        "<AGENT_NAME> Evaluation Baseline",
        "Quality, capability, safety, and regression findings",
        [
            "Agent version: <AGENT_VERSION>",
            "Environment: <ENVIRONMENT> | Run date: <RUN_DATE>",
            "Evaluation owner: <OWNER>",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Executive summary", "Baseline")
    add_metric(slide, 0.65, "SCORED CASES", "<N>", BLUE)
    add_metric(slide, 3.78, "PASS RATE", "<NN%>", GREEN)
    add_metric(slide, 6.91, "BLOCKED", "<N>", AMBER)
    add_metric(slide, 10.04, "CRITICAL FAILURES", "<N>", RED)
    add_text(slide, "Release recommendation", 0.65, 3.52, 3.2, 0.4, size=18, color=NAVY, bold=True)
    add_rect(slide, 0.65, 4.03, 12.0, 1.65, LIGHT_BLUE, LIGHT_BLUE, radius=True)
    add_text(
        slide,
        "<APPROVE | APPROVE WITH CONDITIONS | DO NOT APPROVE>\n"
        "<One-paragraph decision rationale and the most important evidence.>",
        0.98,
        4.35,
        11.35,
        1.0,
        size=18,
        color=DARK,
    )
    add_footer(slide, len(prs.slides))

    add_two_column_slide(
        prs,
        "Scope and run configuration",
        "Baseline",
        "Evaluation scope",
        [
            "Purpose: <release gate / baseline / regression>",
            "Dataset: <NAME_AND_VERSION>",
            "Designed cases: <N>; observed-pattern cases: <N>",
            "In-scope capabilities: <LIST>",
            "Out-of-scope behavior: <LIST>",
        ],
        "Run configuration",
        [
            "Target: <built-in Evaluation / Kit / API>",
            "Agent state: <draft / published>",
            "User profile: <PROFILE>",
            "Evaluation methods: <METHODS>",
            "Thresholds and judge version: <DETAILS>",
        ],
    )

    add_two_column_slide(
        prs,
        "Coverage",
        "Results",
        "Coverage dimensions",
        [
            "Capabilities: <N_OF_N>",
            "Tools and actions: <N_OF_N>",
            "Multi-turn and recovery flows: <N>",
            "Negative and safety scenarios: <N>",
            "Locales, channels, and personas: <DETAILS>",
        ],
        "Coverage gaps",
        [
            "<Missing capability or tool>",
            "<Unrepresented user segment>",
            "<Untested integration failure>",
            "<Known telemetry limitation>",
            "<Planned follow-up date and owner>",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Results by category", "Results")
    add_text(slide, "Replace this placeholder with a chart generated from normalized results.", 0.9, 1.75, 11.5, 0.45, size=17, color=MID_GREY, align=PP_ALIGN.CENTER)
    add_rect(slide, 1.15, 2.45, 11.0, 2.55, LIGHT_GREY, MID_GREY, radius=True)
    add_text(
        slide,
        "<CATEGORY CHART>\n\n"
        "Recommended: pass / fail / blocked counts by capability, risk, tool, or source.",
        1.45,
        3.03,
        10.4,
        1.15,
        size=20,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(slide, "Denominator: <DEFINE WHICH STATUSES COUNT> | Baseline comparison: <CHANGE>", 1.0, 5.48, 11.3, 0.4, size=14, color=DARK, align=PP_ALIGN.CENTER)
    add_footer(slide, len(prs.slides))

    add_two_column_slide(
        prs,
        "Failure analysis",
        "Findings",
        "Failure themes",
        [
            "<Theme 1>: <count and impact>",
            "<Theme 2>: <count and impact>",
            "<Theme 3>: <count and impact>",
            "Evaluator or test-data defects: <N>",
            "Environment or permission blocks: <N>",
        ],
        "Root causes and evidence",
        [
            "<Agent behavior or orchestration cause>",
            "<Tool selection or parameter cause>",
            "<Knowledge or grounding cause>",
            "<Safety or authorization cause>",
            "Use synthetic excerpts only in public reports",
        ],
    )

    add_two_column_slide(
        prs,
        "Actions and release decision",
        "Decision",
        "Required actions",
        [
            "<Action, owner, priority, due date>",
            "<Action, owner, priority, due date>",
            "<Regression case to add>",
            "<Evaluation defect to correct separately>",
            "<Blocked dependency to resolve>",
        ],
        "Decision record",
        [
            "Decision: <STATUS>",
            "Approver: <NAME_OR_ROLE>",
            "Conditions: <LIST>",
            "Next full run: <DATE_OR_TRIGGER>",
            "Baseline artifact location: <URI_OR_PATH>",
        ],
    )

    add_two_column_slide(
        prs,
        "Run metadata and limitations",
        "Appendix",
        "Reproducibility metadata",
        [
            "Run ID: <RUN_ID>",
            "Dataset hash/version: <VALUE>",
            "Solution/model version: <VALUE>",
            "Environment and region: <VALUE>",
            "Results export: <URI_OR_PATH>",
        ],
        "Known limitations",
        [
            "<Judge limitation>",
            "<Telemetry or tool-correlation limitation>",
            "<Missing production segment>",
            "<Non-deterministic dependency>",
            "<Retention or privacy constraint>",
        ],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def build_observability_template(path: Path):
    prs = new_deck(
        "Agent Observability Walkthrough Template",
        "Generic Copilot Studio observability walkthrough template",
    )
    add_title_slide(
        prs,
        "<AGENT_NAME> Observability Walkthrough",
        "Operational health, tool performance, and conversation troubleshooting",
        [
            "Environment: <ENVIRONMENT>",
            "Telemetry resource: <APPLICATION_INSIGHTS_RESOURCE>",
            "Presenter: <OWNER> | Date: <DATE>",
        ],
    )

    add_two_column_slide(
        prs,
        "Objectives and audience",
        "Overview",
        "Questions this walkthrough answers",
        [
            "Is the agent available and completing requests?",
            "Which tools run most often and fail most often?",
            "Where is latency concentrated?",
            "Can support trace a user or conversation?",
            "Which signals are unavailable or approximate?",
        ],
        "Primary audiences",
        [
            "Agent product and engineering owners",
            "Operations and support teams",
            "Platform and integration teams",
            "Security and privacy reviewers",
            "Release and service owners",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Telemetry architecture", "Architecture")
    labels = [
        (0.7, "Copilot Studio\nand integrations", LIGHT_BLUE),
        (3.55, "Application Insights\ntelemetry", LIGHT_GREY),
        (6.4, "KQL normalization\nand analysis", LIGHT_BLUE),
        (9.25, "Workbooks, alerts,\nand exports", LIGHT_GREY),
    ]
    for index, (x, label, fill) in enumerate(labels):
        add_rect(slide, x, 2.35, 2.35, 1.35, fill, BLUE, radius=True)
        add_text(slide, label, x + 0.12, 2.68, 2.1, 0.68, size=17, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        if index < len(labels) - 1:
            add_text(slide, "→", x + 2.42, 2.7, 0.35, 0.4, size=24, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        "Document exact event names, dimensions, retention, and correlation confidence for this environment.",
        1.0,
        4.52,
        11.3,
        0.7,
        size=18,
        color=DARK,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, len(prs.slides))

    add_two_column_slide(
        prs,
        "Operational health dashboard",
        "Dashboard tour",
        "Show",
        [
            "Request and conversation volume",
            "Success, failure, and error trends",
            "Latency p50, p95, and p99",
            "Top tools by calls and failures",
            "Trend comparison with the prior period",
        ],
        "Explain",
        [
            "Selected time range and agent filters",
            "Definitions of calls, failures, and latency",
            "Alert or service-level thresholds",
            "Expected traffic and known seasonality",
            "Data freshness and retention",
        ],
    )

    add_two_column_slide(
        prs,
        "Conversation troubleshooting",
        "Dashboard tour",
        "Investigation workflow",
        [
            "Find the user or conversation",
            "Review ordered user and agent messages",
            "Overlay tool calls and durations",
            "Identify errors, retries, and long waits",
            "Capture evidence without exposing sensitive text",
        ],
        "Correlation caveats",
        [
            "State which joins are exact",
            "Label timestamp-window matches as approximate",
            "Account for overlapping conversations",
            "Distinguish tool latency from end-to-end latency",
            "Record missing identifiers as telemetry gaps",
        ],
    )

    add_two_column_slide(
        prs,
        "Cost, tokens, and capacity",
        "Limitations",
        "Available signals",
        [
            "<Token usage by model, if emitted>",
            "<Capacity or billed-message metrics>",
            "<Owned model or integration costs>",
            "<Request-volume proxy>",
            "<Current pricing source and date>",
        ],
        "Limitations to disclose",
        [
            "Copilot Studio may not emit token counts",
            "Pricing is configuration, not telemetry",
            "Regional and provisioned pricing can differ",
            "Conversation IDs may vary by source",
            "Never report unsupported precision",
        ],
    )

    add_two_column_slide(
        prs,
        "Operational response",
        "Operations",
        "When a metric breaches",
        [
            "Confirm scope, time range, and data freshness",
            "Segment by agent, tool, channel, and version",
            "Inspect representative failed conversations",
            "Check dependent services and permissions",
            "Open an incident with reproducible evidence",
        ],
        "Ownership",
        [
            "Agent behavior: <OWNER>",
            "Tool or integration: <OWNER>",
            "Platform and telemetry: <OWNER>",
            "Security and privacy: <OWNER>",
            "Escalation path: <PROCESS_OR_LINK>",
        ],
    )

    add_two_column_slide(
        prs,
        "Rollout checklist",
        "Next steps",
        "Before launch",
        [
            "Verify telemetry schema with discovery queries",
            "Set retention and access controls",
            "Publish workbooks and saved queries",
            "Define thresholds and alert ownership",
            "Test a known synthetic conversation",
        ],
        "Ongoing",
        [
            "Review health and quality together",
            "Revalidate queries after platform changes",
            "Update pricing and thresholds",
            "Promote incident patterns into evaluations",
            "Remove sensitive exports when retention ends",
        ],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        default="reports",
        help="Directory for generated template decks (default: reports)",
    )
    args = parser.parse_args()

    output_dir = Path(args.output_dir)
    baseline = output_dir / "baseline-evaluation-report-template.pptx"
    observability = output_dir / "observability-walkthrough-template.pptx"
    build_baseline_template(baseline)
    build_observability_template(observability)
    print(f"Wrote {baseline}")
    print(f"Wrote {observability}")


if __name__ == "__main__":
    main()
